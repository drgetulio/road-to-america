#!/usr/bin/env python3
"""Planilha de produção da Questão de Embarque a partir de instagram/programa-365.json.

    python3 instagram/ferramentas/gerar-planilha.py

Saída: entregas/instagram-365-dias.xlsx com as abas Calendário, Prompts das imagens,
Personagens e Como usar.
"""
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
PROG = json.loads((RAIZ / 'instagram' / 'programa-365.json').read_text(encoding='utf8'))
SAIDA = RAIZ / 'entregas' / 'instagram-365-dias.xlsx'

NAVY, VERMELHO, PAPEL, VERDE, AMARELO = '0F2747', 'B0243A', 'F3F5F2', '1D8A4B', 'F1C232'
F = 'Arial'
fonte = Font(name=F, size=10)
cab_font = Font(name=F, size=10, bold=True, color='FFFFFF')
cab_fill = PatternFill('solid', fgColor=NAVY)
status_fill = PatternFill('solid', fgColor='FFF6D6')
remix_fill = PatternFill('solid', fgColor='E7EBE6')
acolhe_fill = PatternFill('solid', fgColor='E6EEF8')
fino = Side(style='thin', color='D0D6D0')
borda = Border(left=fino, right=fino, top=fino, bottom=fino)
quebra = Alignment(wrap_text=True, vertical='top')
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']

STATUS_IMG = ['Pendente', 'Gerada', 'Aprovada', 'Refazer']
STATUS_RENDER = ['Pendente', 'Renderizado', 'Conferido']
STATUS_POST = ['Pendente', 'Agendado', 'Postado', 'Adiado']
STATUS_REV = ['Pendente', 'Revisado', 'Ajustar']


def data_br(iso):
    a, m, d = iso.split('-')
    return f'{d}/{m}/{a}'


def cabecalho(ws, colunas, linha=1):
    for i, (nome, larg) in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.font = cab_font
        c.fill = cab_fill
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = borda
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[linha].height = 42


def escrever(ws, linha, valores, fill=None):
    for i, v in enumerate(valores, start=1):
        c = ws.cell(row=linha, column=i, value=v)
        c.font = fonte
        c.alignment = quebra
        c.border = borda
        if fill:
            c.fill = fill


def validacao(ws, opcoes, col, ini, fim):
    dv = DataValidation(type='list', formula1='"' + ','.join(opcoes) + '"', allow_blank=True)
    dv.error = 'Escolha um dos valores da lista.'
    dv.prompt = ' / '.join(opcoes)
    ws.add_data_validation(dv)
    dv.add(f'{col}{ini}:{col}{fim}')


def letra(n):
    return get_column_letter(n)


wb = Workbook()

# ------------------------------------------------------------------ Calendário
ws = wb.active
ws.title = 'Calendário'
COLS = [
    ('Dia', 6), ('Data (D)', 11), ('Dia da semana', 12), ('Mês', 8), ('Quadro', 16), ('Step', 9), ('Sistema', 22), ('Área', 18),
    ('Tópico', 26), ('Conceito (inglês)', 34), ('Tipo de pergunta', 13), ('Personagem', 14), ('Gancho de data', 26), ('Remix de (dias)', 11),
    ('Tom', 11), ('CTA do dia', 12), ('Gancho (capa)', 34), ('Gancho de apoio', 28), ('Vinheta (inglês)', 60), ('Pergunta', 30),
    ('A', 24), ('B', 24), ('C', 24), ('D', 24), ('Correta', 8), ('Resposta certa', 28), ('Por que a certa está certa', 60),
    ('Por que as erradas estão erradas', 60), ('Termos da resposta (só em D+1)', 30), ('Macete', 40), ('Pegadinha de prova', 40),
    ('Palavras-chave neutras (legenda D)', 30), ('Legenda D (só a questão)', 60), ('Legenda do post de D (questão D + gabarito D-1, pronta para colar)', 70),
    ('Legenda D+1 (gabarito e explicação, comentário fixado)', 70), ('Explicação sai em (D+1)', 11), ('Story: pergunta do quiz', 30),
    ('Story: opção A', 16), ('Story: opção B', 16), ('Story: opção C', 16), ('Story: opção D', 16), ('Story: letra certa', 8),
    ('Roteiro dos stories', 60), ('Imagem da questão (arquivo)', 26), ('Imagem da explicação (arquivo)', 28), ('Fonte', 40),
    ('Revisão médica', 50), ('Revisão de humor e compliance', 50),
    ('Status imagem questão', 13), ('Status imagem explicação', 13), ('Status render', 13), ('Revisão final (fontes)', 13),
    ('Status feed (D)', 12), ('Status story quiz (D)', 12), ('Status gabarito (D+1)', 12), ('Postado em', 11), ('Alcance', 10),
    ('Salvamentos', 10), ('Taxa de salvamento', 10), ('Respostas no quiz', 10), ('Acerto no quiz (%)', 10), ('Observações', 30),
]
cabecalho(ws, COLS)
idx = {nome: i + 1 for i, (nome, _) in enumerate(COLS)}
N = len(PROG['dias'])
for r, d in enumerate(PROG['dias'], start=2):
    q, ex, lg, st, im = d['questao'], d['explicacao'], d['legendas'], d['story'], d['imagens']
    alts = {a['letra']: a for a in q['alternativas']}
    altx = lambda l: alts[l]['texto'] + (f" (visto no dia {alts[l]['visto_no_dia']})" if alts[l]['visto_no_dia'] else '')
    erradas = '\n'.join(f'{k}) {v}' for k, v in ex['erradas'].items())
    rev = d['revisao'] or {}
    mes = d['data'][5:7]
    vals = [
        d['dia'], data_br(d['data']), d['dia_semana'], f"{MESES[int(mes) - 1]}/{d['data'][2:4]}", d['quadro'], d['step'], d['sistema'], d['area'],
        d['topico'], d['conceito'], d['tipo_pergunta'], d['personagem'] or '', d['gancho_data'] or '',
        ', '.join(map(str, d['remix_de'])) if d['remix_de'] else '', d['tom'], d['cta'], d['gancho'], d['gancho_apoio'] or '',
        q['vinheta'], q['pergunta'], altx('A'), altx('B'), altx('C'), altx('D'), q['correta'], q['resposta'], ex['certa'], erradas,
        ', '.join(ex['termos']), ex['macete'] + (f"\n{ex['macete_apoio']}" if ex['macete_apoio'] else ''), ex['pegadinha'] or '',
        ', '.join(lg['palavras_chave']), lg['questao_D'], lg['post_D'], lg['explicacao_D1'], data_br(d['explicacao_em']), st['pergunta'],
        *st['opcoes'], st['correta'], st['roteiro'], im['questao']['arquivo'], im['explicacao']['arquivo'], d['fonte'],
        (rev.get('medica', '') + (f"\n{rev['letras']}" if rev.get('letras') else '')).strip(), rev.get('humor_compliance', ''),
        'Pendente', 'Pendente', 'Pendente', 'Pendente', 'Pendente', 'Pendente', 'Pendente', None, None, None, None, None, None, '',
    ]
    fill = remix_fill if d['remix_de'] else acolhe_fill if d['tom'] == 'acolhedor' else None
    escrever(ws, r, vals, fill)
    ws.row_dimensions[r].height = 96
    # taxa de salvamento = salvamentos / alcance
    a, s = letra(idx['Alcance']), letra(idx['Salvamentos'])
    c = ws.cell(row=r, column=idx['Taxa de salvamento'], value=f'=IF(AND(ISNUMBER({a}{r}),{a}{r}>0,ISNUMBER({s}{r})),{s}{r}/{a}{r},"")')
    c.number_format = '0.0%'
    ws.cell(row=r, column=idx['Postado em']).number_format = 'DD/MM/YYYY'
    for nome in ['Status imagem questão', 'Status imagem explicação', 'Status render', 'Revisão final (fontes)', 'Status feed (D)', 'Status story quiz (D)', 'Status gabarito (D+1)', 'Postado em', 'Alcance', 'Salvamentos', 'Respostas no quiz', 'Acerto no quiz (%)', 'Observações']:
        ws.cell(row=r, column=idx[nome]).fill = status_fill
    ws.cell(row=r, column=idx['Correta']).font = Font(name=F, size=11, bold=True, color=VERMELHO)
    ws.cell(row=r, column=idx['Correta']).alignment = Alignment(horizontal='center', vertical='top')
ult = N + 1
for nome in ['Status imagem questão', 'Status imagem explicação']:
    validacao(ws, STATUS_IMG, letra(idx[nome]), 2, ult)
validacao(ws, STATUS_RENDER, letra(idx['Status render']), 2, ult)
validacao(ws, STATUS_REV, letra(idx['Revisão final (fontes)']), 2, ult)
for nome in ['Status feed (D)', 'Status story quiz (D)', 'Status gabarito (D+1)']:
    validacao(ws, STATUS_POST, letra(idx[nome]), 2, ult)
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{letra(len(COLS))}{ult}'

# ------------------------------------------------------------------ Prompts das imagens
wp = wb.create_sheet('Prompts das imagens')
PCOLS = [('Dia', 6), ('Data de uso', 11), ('Tipo', 11), ('Quadro', 16), ('Personagem', 14), ('Prompt pronto para copiar (Higgsfield)', 90),
         ('Modelo', 16), ('Resolução', 9), ('Proporção', 9), ('Nome de arquivo esperado', 30), ('Cena em português', 50), ('Texto alternativo', 34),
         ('Status da geração', 13), ('Observações', 30)]
cabecalho(wp, PCOLS)
r = 2
for d in PROG['dias']:
    for tipo, chave, quando in [('questão', 'questao', d['data']), ('explicação', 'explicacao', d['explicacao_em'])]:
        im = d['imagens'][chave]
        escrever(wp, r, [d['dia'], data_br(quando), tipo, d['quadro'], d['personagem'] or '', im['prompt'], 'Nano Banana Pro', '2k', '4:5',
                         im['arquivo'], im['cena'], im['alt'] or '', 'Pendente', ''], remix_fill if tipo == 'explicação' else None)
        wp.cell(row=r, column=13).fill = status_fill
        wp.cell(row=r, column=14).fill = status_fill
        wp.row_dimensions[r].height = 150
        r += 1
validacao(wp, STATUS_IMG, 'M', 2, r - 1)
wp.freeze_panes = 'D2'
wp.auto_filter.ref = f'A1:{letra(len(PCOLS))}{r - 1}'

# ------------------------------------------------------------------ Personagens
wc = wb.create_sheet('Personagens')
CCOLS = [('Personagem', 18), ('Quem é', 40), ('Descrição visual fixa (português)', 50), ('Descrição fixa para o prompt (inglês)', 60),
         ('Prompt-base (acrescente a cena do dia)', 90), ('Bordão', 26), ('Onde aparece', 30), ('Dias em que é o personagem do dia', 18), ('Lista dos dias', 40)]
cabecalho(wc, CCOLS)
r = 2
for p in PROG['personagens'] + PROG['apoio']:
    escrever(wc, r, [p['nome'], p['quem'], p['pt'], p['descricao_en'], p['prompt_base'], p['bordao'] or '', p['onde'], len(p['dias']) or '',
                     ', '.join(map(str, p['dias']))])
    wc.row_dimensions[r].height = 170
    r += 1
wc.freeze_panes = 'B2'
wc.auto_filter.ref = f'A1:{letra(len(CCOLS))}{r - 1}'

# ------------------------------------------------------------------ Como usar
wu = wb.create_sheet('Como usar')
wu.column_dimensions['A'].width = 34
wu.column_dimensions['B'].width = 70
wu.column_dimensions['C'].width = 26
wu.column_dimensions['D'].width = 50
titulo = Font(name=F, size=14, bold=True, color=NAVY)
sub = Font(name=F, size=11, bold=True, color=VERMELHO)
linha = [1]


def bloco(t):
    linha[0] += 1
    c = wu.cell(row=linha[0], column=1, value=t)
    c.font = sub
    linha[0] += 1


def item(a, b='', c='', d=''):
    for i, v in enumerate([a, b, c, d], start=1):
        x = wu.cell(row=linha[0], column=i, value=v)
        x.font = fonte
        x.alignment = quebra
    linha[0] += 1


wu.cell(row=1, column=1, value='Questão de Embarque: como usar esta planilha').font = titulo
linha[0] = 2
item('Programa', f"365 dias, de {data_br(PROG['periodo']['inicio'])} a {data_br(PROG['periodo']['fim'])}; o gabarito do dia 365 sai em {data_br(PROG['periodo']['ultima_explicacao'])}, no post de encerramento.")
nums = PROG['numeros']
item('Números', f"{nums['dias']} questões ({nums['por_step']['Step 1']} de Step 1 e {nums['por_step']['Step 2 CK']} de Step 2 CK), {nums['remix']} remixes de domingo, {nums['imagens']} imagens, {nums['tom_acolhedor']} dias em tom acolhedor. Letra certa: A {nums['letras']['A']}, B {nums['letras']['B']}, C {nums['letras']['C']}, D {nums['letras']['D']}.")
item('Cores das linhas', 'Cinza: Remix de Domingo. Azul claro: data sensível, tom acolhedor e sem piada. Amarelo claro: colunas de status para preencher (lista suspensa).')
item('Onde editar', 'Só as colunas amarelas (status, data de postagem, métricas e observações). O conteúdo vem de instagram/dias/; corrija lá e gere a planilha de novo.')
item('Exemplo de preenchimento', 'Status imagem questão = Aprovada; Status feed (D) = Postado; Postado em = 02/11/2026; Alcance = 4200; Salvamentos = 310 (a taxa de salvamento se calcula sozinha); Acerto no quiz (%) = 58.')

bloco('Rotina diária (horário de Brasília)')
for h, o in [('11h50', 'Story "Gabarito em 10 minutos" com a imagem do macete de ontem desfocada.'),
             ('12h00', 'Feed: carrossel de 8 slides (3 da questão de hoje + 5 do gabarito de ontem) com a coluna "Legenda do post de D". No dia 1, só os 3 slides da questão.'),
             ('12h01', 'Primeiro comentário fixado no post: a "Legenda D+1" do dia anterior (explicação completa).'),
             ('12h05', 'Story do quiz: story-quiz.png + adesivo de QUIZ com a pergunta e as quatro opções curtas da planilha; marcar a letra certa.'),
             ('12h06', 'Enquete "Você sabia ou chutou?" (Sabia / Chutei).'),
             ('20h30', 'Resultado parcial sem revelar a letra + adesivo de link para o post.'),
             ('Domingo 18h00', 'Placar da semana: as três questões com menor acerto + link do simulado grátis.'),
             ('Primeiras 4 semanas', 'Testar 20h30 nas terças e quintas e ficar com o horário de mais salvamentos por alcance.')]:
    item(h, o)

bloco('Stories com quiz e CTA')
item('Quiz', 'O adesivo aceita no máximo 4 opções; as opções curtas (até 30 caracteres) estão nas colunas Story: opção A a D. O quiz mostra a certa a quem responde, por isso a explicação completa só sai no D+1.')
item('CTA', 'Terça, quinta e domingo: simulado grátis. Domingo: pré-inscrição. Primeira quarta do mês: verificador de faculdade. Outros dias: salvar e mandar para quem vai fazer o Step. Links só na bio e no adesivo de link, sempre com UTM (o roteiro dos stories traz os links prontos).')
item('Destaques fixos', 'Step 1, Step 2 CK, Remix, Simulado grátis, Pré-inscrição.')

bloco('Fluxo de geração das imagens (Higgsfield)')
for a, b in [('1. Abrir', 'Higgsfield, gerador de imagem, modelo Nano Banana Pro, resolução 2k, proporção 4:5.'),
             ('2. Copiar', 'Aba "Prompts das imagens": copie a célula do prompt inteira (estilo, personagem fixo, cena, composição e proibições já estão nela). Duas imagens por dia: questão (D) e explicação (D+1).'),
             ('3. Conferir', 'Sem texto, letras ou números na imagem; personagem igual à aba Personagens; terço de baixo liso (a placa do gancho cobre essa área); a imagem da questão não pode mostrar a resposta.'),
             ('4. Salvar', 'Com o nome da coluna "Nome de arquivo esperado": instagram/imagens/dia-NNN.png e dia-NNN-macete.png.'),
             ('5. Marcar', 'Status da geração = Gerada; depois da revisão, Aprovada (ou Refazer, com o motivo em Observações).'),
             ('Ritmo', 'Gere em lotes de uma semana, com pelo menos 7 dias de antecedência.')]:
    item(a, b)

bloco('Renderização dos posts')
for a, b in [('Validar o ano', 'node instagram/ferramentas/validar-dias.mjs (grava instagram/questoes/ e roda render.mjs --so-conferir nos 365 dias, com a comparação com o banco privado).'),
             ('Renderizar uma semana', 'node instagram/templates/render.mjs --intervalo 15-21'),
             ('Saída', 'instagram/saida/dia-NNN/: post-do-dia/ (carrossel numerado), legenda.txt, stories.txt, story-quiz.png, story-quiz-guia.png e manifesto.json.'),
             ('Sem imagem', 'Os slides saem com a marca IMAGEM PENDENTE e o nome do arquivo esperado. Gere a imagem e renderize de novo.'),
             ('Atualizar a planilha', 'node instagram/ferramentas/montar-programa.mjs && python3 instagram/ferramentas/gerar-planilha.py')]:
    item(a, b)

bloco('Regras de humor e compliance (resumo)')
for a, b in [('Humor', 'A piada aponta para a resposta e continuaria engraçada se alguém com a doença estivesse lendo. Nunca sobre paciente, corpo, grupos, profissões, doença grave ou estigmatizada.'),
             ('Datas sensíveis', 'Finados, AIDS, Janeiro Branco, Câncer, Down, Autismo, Setembro Amarelo (CVV 188), Alzheimer, Saúde Mental, início do Outubro Rosa e do Novembro Azul: personagens explicam, sem gag.'),
             ('CFM 2.336/2023', 'Sem promessa de aprovação, nota ou Match; sem "o melhor curso"; nome e CRM-DF 13610 em toda legenda; especialidade só com RQE.'),
             ('Imagens', 'Só ilustração própria; nada de paciente, exame real, pessoa reconhecível ou material de terceiros.')]:
    item(a, b)

bloco('Fontes que precisam ser reconferidas antes de postar (compiladas dos relatórios dos revisores)')
for i, v in enumerate(['Dia', 'Data', 'O que conferir', 'Fonte citada']):
    c = wu.cell(row=linha[0], column=i + 1, value=v)
    c.font = cab_font
    c.fill = cab_fill
linha[0] += 1
pend = re.compile(r'(re)?conferir|estiver vigente|vers[aã]o vigente|edi[cç][aã]o vigente', re.I)
ignorar = re.compile(r'gancho dizia|^IRT', re.I)
n_fontes = 0
for d in PROG['dias']:
    rev = d['revisao'] or {}
    frases = []
    for k in ('medica', 'humor_compliance'):
        for s in re.split(r'(?<=[.;])\s+', rev.get(k, '')):
            if pend.search(s) and not ignorar.search(s) and not re.match(r'^Conferid', s):
                frases.append(s.strip().rstrip(';'))
    if frases:
        item(f"Dia {d['dia']} ({d['quadro']})", data_br(d['data']), ' '.join(frases), d['fonte'])
        n_fontes += 1
item('Datas móveis de 2027', 'Carnaval 09/02, Cinzas 10/02, Sexta-feira Santa 26/03, Páscoa 28/03, Match Day 19/03 (NRMP) e prazo do rank order list 03/03: reconferir uma semana antes de cada post.')
item('Todo trimestre', 'Questões de rastreamento, vacinas e conduta: conferir a diretriz vigente (lista por instituição abaixo).')

bloco('Diretrizes citadas, por instituição (dias)')
orgs = ['CDC', 'ACIP', 'USPSTF', 'Preventive Services Task Force', 'AHA', 'ACC', 'AAP', 'ACOG', 'IDSA', 'ADA', 'GINA', 'GOLD', 'KDIGO', 'AASLD', 'ACG', 'AGA',
        'ACR', 'AUA', 'ESC', 'ATLS', 'Endocrine Society', 'Joint Commission', 'AMA', 'FDA', 'WHO', 'ASA', 'ATA', 'AAOS', 'ACCP', 'AAN', 'NCCN', 'Surviving Sepsis']
grupo = defaultdict(list)
for d in PROG['dias']:
    for o in orgs:
        if re.search(r'\b' + re.escape(o) + r'\b', d['fonte'] or ''):
            grupo['USPSTF' if o == 'Preventive Services Task Force' else o].append(d['dia'])
for o, ds in sorted(grupo.items(), key=lambda x: -len(x[1])):
    item(o, ', '.join(map(str, sorted(set(ds)))), f'{len(set(ds))} dias')

bloco('Andamento (atualiza sozinho com os status da aba Calendário)')
cal_ref = "'Calendário'!"
for nome, opcoes in [('Status imagem questão', STATUS_IMG), ('Status imagem explicação', STATUS_IMG), ('Status render', STATUS_RENDER), ('Status feed (D)', STATUS_POST)]:
    col = letra(idx[nome])
    for o in opcoes:
        wu.cell(row=linha[0], column=1, value=f'{nome}: {o}').font = fonte
        c = wu.cell(row=linha[0], column=2, value=f'=COUNTIF({cal_ref}{col}2:{col}{ult},"{o}")')
        c.font = fonte
        c.alignment = Alignment(horizontal='left')
        linha[0] += 1
col_t = letra(idx['Taxa de salvamento'])
wu.cell(row=linha[0], column=1, value='Taxa média de salvamento').font = fonte
c = wu.cell(row=linha[0], column=2, value=f'=IFERROR(AVERAGE({cal_ref}{col_t}2:{col_t}{ult}),"")')
c.number_format = '0.0%'
c.alignment = Alignment(horizontal='left')
linha[0] += 1


wb.active = 0
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print(f'{SAIDA.relative_to(RAIZ)}: {N} dias, {2 * N} prompts, {len(PROG["personagens"]) + len(PROG["apoio"])} personagens e apoios, {n_fontes} dias com fonte a reconferir')
